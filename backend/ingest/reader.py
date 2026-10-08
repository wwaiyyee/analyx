"""Raw file ingestion: format detection, encoding, delimiter sniffing, and XLSX reading."""

import csv
import io
from pathlib import Path
from typing import Any, Tuple

import chardet
import openpyxl

from backend.core.errors import IngestError
from backend.core.ids import generate_id
from backend.ingest.models import TransformationStep

MAX_FILE_BYTES = 50 * 1024 * 1024  # 50 MB
CANDIDATE_DELIMITERS = [",", ";", "\t", "|"]


def detect_file_format(filename: str, content: bytes) -> str:
    """Detect file format using extension and magic bytes."""
    ext = Path(filename).suffix.lower()

    if content.startswith(b"PK\x03\x04"):
        if ext in (".xlsx", ".xlsm", ".xltx", ".xltm"):
            return "xlsx"
        # ZIP archive but not xlsx extension
        return "xlsx"

    if ext == ".xlsx":
        return "xlsx"

    if ext in (".csv", ".tsv", ".txt") or not ext:
        # Verify it doesn't look like an executable or raw binary
        if content.startswith(b"\x7fELF") or content.startswith(b"MZ"):
            raise IngestError(f"Unsupported binary executable file: {filename}")
        return "csv"

    raise IngestError(f"Unsupported file format '{ext}' for file '{filename}'. Supported: .csv, .xlsx")


def detect_csv_encoding(content: bytes) -> str:
    """Detect text encoding, preferring UTF-8 / UTF-8-SIG."""
    if content.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    if content.startswith(b"\xff\xfe") or content.startswith(b"\xfe\xff"):
        return "utf-16"

    # Try fast utf-8 decode
    try:
        content[:65536].decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError:
        pass

    # Chardet fallback
    detected = chardet.detect(content[:65536])
    encoding = detected.get("encoding")
    if encoding and detected.get("confidence", 0) > 0.6:
        return encoding
    return "latin-1"


def detect_csv_delimiter(sample_text: str) -> str:
    """Detect the most likely delimiter from candidate set using line consistency."""
    lines = [line for line in sample_text.splitlines()[:30] if line.strip()]
    if not lines:
        return ","

    try:
        sniffer = csv.Sniffer()
        dialect = sniffer.sniff(sample_text[:4096], delimiters=",\t;|")
        if dialect.delimiter in CANDIDATE_DELIMITERS:
            return dialect.delimiter
    except Exception:
        pass

    # Fallback: Count frequency and variance of candidates across lines
    best_delim = ","
    best_score = -1.0

    for delim in CANDIDATE_DELIMITERS:
        counts = [line.count(delim) for line in lines]
        if all(c > 0 for c in counts):
            avg = sum(counts) / len(counts)
            variance = sum((c - avg) ** 2 for c in counts) / len(counts)
            if variance < 0.1:  # Perfectly consistent column count
                score = avg * 100
            else:
                score = avg / (1.0 + variance)
            if score > best_score:
                best_score = score
                best_delim = delim

    return best_delim


def read_csv_raw(content: bytes) -> Tuple[list[list[Any]], dict[str, Any]]:
    """Decode and parse raw CSV content into rows and metadata."""
    if len(content) > MAX_FILE_BYTES:
        raise IngestError(f"File size exceeds limit of {MAX_FILE_BYTES // (1024*1024)} MB")

    encoding = detect_csv_encoding(content)
    text = content.decode(encoding, errors="replace")

    delimiter = detect_csv_delimiter(text[:8192])
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)

    raw_rows: list[list[Any]] = []
    for row in reader:
        raw_rows.append([cell.strip() for cell in row])

    meta = {
        "encoding": encoding,
        "delimiter": delimiter,
        "raw_row_count": len(raw_rows),
    }
    return raw_rows, meta


def read_xlsx_raw(content: bytes) -> Tuple[list[list[Any]], dict[str, Any], list[TransformationStep]]:
    """Read Excel workbook, selecting the largest non-empty sheet."""
    if len(content) > MAX_FILE_BYTES:
        raise IngestError(f"File size exceeds limit of {MAX_FILE_BYTES // (1024*1024)} MB")

    workbook = openpyxl.load_workbook(
        io.BytesIO(content),
        read_only=True,
        data_only=True,  # Read cached formula values, never execute code
    )

    sheet_names = workbook.sheetnames
    if not sheet_names:
        raise IngestError("Workbook contains no sheets")

    # Find the largest non-empty sheet
    largest_sheet_name = sheet_names[0]
    max_cells = -1
    best_rows: list[list[Any]] = []

    for name in sheet_names:
        sheet = workbook[name]
        sheet_rows: list[list[Any]] = []
        for row in sheet.iter_rows(values_only=True):
            cleaned = [str(c).strip() if c is not None else "" for c in row]
            if any(cleaned):
                sheet_rows.append(cleaned)
        cell_count = sum(len(r) for r in sheet_rows)
        if cell_count > max_cells:
            max_cells = cell_count
            largest_sheet_name = name
            best_rows = sheet_rows

    workbook.close()

    transformations: list[TransformationStep] = []
    if len(sheet_names) > 1:
        transformations.append(
            TransformationStep(
                id=generate_id("tr"),
                operation="select_sheet",
                parameters={"sheet": largest_sheet_name, "available_sheets": sheet_names},
                rows_before=len(best_rows),
                rows_after=len(best_rows),
                requires_approval=False,
            )
        )

    meta = {
        "sheet_names": sheet_names,
        "selected_sheet": largest_sheet_name,
        "raw_row_count": len(best_rows),
    }
    return best_rows, meta, transformations


def read_file(filename: str, content: bytes) -> Tuple[list[list[Any]], dict[str, Any], list[TransformationStep]]:
    """Universal raw file reader dispatching to CSV or XLSX parser."""
    fmt = detect_file_format(filename, content)
    if fmt == "xlsx":
        return read_xlsx_raw(content)
    raw_rows, meta = read_csv_raw(content)
    return raw_rows, meta, []
